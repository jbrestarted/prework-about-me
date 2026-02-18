"""
Systems Analysis Word Document Generator

Generates professional .docx documents with a structured systems analysis template.
"""

import io
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT


def _set_cell_shading(cell, color_hex: str):
    """Apply background shading to a table cell."""
    from docx.oxml.ns import qn
    from lxml import etree

    shading = etree.SubElement(cell._element.get_or_add_tcPr(), qn("w:shd"))
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")


def _add_styled_table(doc, headers: list[str], rows: list[list[str]]):
    """Add a formatted table with header row shading."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_cell_shading(cell, "2E4057")

    # Data rows
    for r, row_data in enumerate(rows):
        for c, value in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    return table


def generate_document(data: dict) -> io.BytesIO:
    """
    Generate a Systems Analysis Word document from form data.

    Args:
        data: Dictionary with all form field values.

    Returns:
        BytesIO buffer containing the .docx file.
    """
    doc = Document()

    # -- Page setup --
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    # -- Default font --
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style.paragraph_format.space_after = Pt(6)

    # -- Heading styles --
    for level in range(1, 4):
        h_style = doc.styles[f"Heading {level}"]
        h_style.font.name = "Calibri"
        h_style.font.color.rgb = RGBColor(0x2E, 0x40, 0x57)

    # ========================================================================
    # TITLE PAGE
    # ========================================================================
    for _ in range(6):
        doc.add_paragraph("")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Systems Analysis Document")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x2E, 0x40, 0x57)

    doc.add_paragraph("")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(data.get("project_name", "Untitled Project"))
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph("")

    # Meta info block
    meta_items = [
        ("Prepared by", data.get("author", "")),
        ("Organization", data.get("organization", "")),
        ("Date", data.get("document_date", date.today().isoformat())),
        ("Version", data.get("version", "1.0")),
    ]
    for label, value in meta_items:
        if value:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"{label}: ")
            run.bold = True
            run.font.size = Pt(11)
            p.add_run(value).font.size = Pt(11)

    doc.add_page_break()

    # ========================================================================
    # TABLE OF CONTENTS PLACEHOLDER
    # ========================================================================
    doc.add_heading("Table of Contents", level=1)
    toc_note = doc.add_paragraph(
        "To generate the Table of Contents in Microsoft Word: "
        "right-click this text and select 'Update Field', or press Ctrl+A "
        "then F9."
    )
    toc_note.italic = True

    # Insert a TOC field code
    from docx.oxml.ns import qn

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_char_begin = run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    run._element.append(fld_char_begin)

    run2 = paragraph.add_run()
    instr = run2._element.makeelement(qn("w:instrText"), {})
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run2._element.append(instr)

    run3 = paragraph.add_run()
    fld_char_end = run3._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    run3._element.append(fld_char_end)

    doc.add_page_break()

    # ========================================================================
    # 1. EXECUTIVE SUMMARY
    # ========================================================================
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(data.get("executive_summary", ""))

    # ========================================================================
    # 2. INTRODUCTION
    # ========================================================================
    doc.add_heading("2. Introduction", level=1)

    doc.add_heading("2.1 Purpose", level=2)
    doc.add_paragraph(data.get("purpose", ""))

    doc.add_heading("2.2 Scope", level=2)
    doc.add_paragraph(data.get("scope", ""))

    doc.add_heading("2.3 Definitions & Acronyms", level=2)
    definitions = data.get("definitions", "")
    if definitions.strip():
        for line in definitions.strip().split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")
    else:
        doc.add_paragraph("No definitions provided.")

    # ========================================================================
    # 3. CURRENT SYSTEM OVERVIEW
    # ========================================================================
    doc.add_heading("3. Current System Overview", level=1)

    doc.add_heading("3.1 Current System Description", level=2)
    doc.add_paragraph(data.get("current_system_description", ""))

    doc.add_heading("3.2 Current System Limitations", level=2)
    limitations = data.get("current_system_limitations", "")
    if limitations.strip():
        for line in limitations.strip().split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")
    else:
        doc.add_paragraph("No limitations documented.")

    # ========================================================================
    # 4. PROPOSED SYSTEM
    # ========================================================================
    doc.add_heading("4. Proposed System", level=1)

    doc.add_heading("4.1 System Overview", level=2)
    doc.add_paragraph(data.get("proposed_system_overview", ""))

    doc.add_heading("4.2 System Goals & Objectives", level=2)
    goals = data.get("system_goals", "")
    if goals.strip():
        for line in goals.strip().split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")
    else:
        doc.add_paragraph("No goals specified.")

    # ========================================================================
    # 5. REQUIREMENTS
    # ========================================================================
    doc.add_heading("5. Requirements", level=1)

    doc.add_heading("5.1 Functional Requirements", level=2)
    func_reqs = data.get("functional_requirements", "")
    if func_reqs.strip():
        _add_requirements_list(doc, func_reqs)
    else:
        doc.add_paragraph("No functional requirements specified.")

    doc.add_heading("5.2 Non-Functional Requirements", level=2)
    nonfunc_reqs = data.get("nonfunctional_requirements", "")
    if nonfunc_reqs.strip():
        _add_requirements_list(doc, nonfunc_reqs)
    else:
        doc.add_paragraph("No non-functional requirements specified.")

    # ========================================================================
    # 6. SYSTEM ARCHITECTURE
    # ========================================================================
    doc.add_heading("6. System Architecture", level=1)

    doc.add_heading("6.1 Architecture Overview", level=2)
    doc.add_paragraph(data.get("architecture_overview", ""))

    doc.add_heading("6.2 Technology Stack", level=2)
    tech_stack = data.get("technology_stack", "")
    if tech_stack.strip():
        for line in tech_stack.strip().split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")
    else:
        doc.add_paragraph("No technology stack specified.")

    doc.add_heading("6.3 Data Flow Description", level=2)
    doc.add_paragraph(data.get("data_flow", ""))

    doc.add_heading("6.4 Integration Points", level=2)
    doc.add_paragraph(data.get("integration_points", ""))

    # ========================================================================
    # 7. DATA REQUIREMENTS
    # ========================================================================
    doc.add_heading("7. Data Requirements", level=1)

    doc.add_heading("7.1 Data Entities", level=2)
    doc.add_paragraph(data.get("data_entities", ""))

    doc.add_heading("7.2 Data Storage & Retention", level=2)
    doc.add_paragraph(data.get("data_storage", ""))

    # ========================================================================
    # 8. USER INTERFACE REQUIREMENTS
    # ========================================================================
    doc.add_heading("8. User Interface Requirements", level=1)
    doc.add_paragraph(data.get("ui_requirements", ""))

    # ========================================================================
    # 9. SECURITY REQUIREMENTS
    # ========================================================================
    doc.add_heading("9. Security Requirements", level=1)
    doc.add_paragraph(data.get("security_requirements", ""))

    # ========================================================================
    # 10. STAKEHOLDER ANALYSIS
    # ========================================================================
    doc.add_heading("10. Stakeholder Analysis", level=1)
    stakeholders = data.get("stakeholders", "")
    if stakeholders.strip():
        rows = []
        for line in stakeholders.strip().split("\n"):
            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 3:
                rows.append(parts[:3])
            elif line.strip():
                rows.append([line.strip(), "", ""])
        if rows:
            _add_styled_table(doc, ["Stakeholder", "Role", "Interest"], rows)
        else:
            doc.add_paragraph(stakeholders)
    else:
        doc.add_paragraph("No stakeholders identified.")

    # ========================================================================
    # 11. RISK ASSESSMENT
    # ========================================================================
    doc.add_heading("11. Risk Assessment", level=1)
    risks = data.get("risks", "")
    if risks.strip():
        rows = []
        for line in risks.strip().split("\n"):
            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 4:
                rows.append(parts[:4])
            elif line.strip():
                rows.append([line.strip(), "", "", ""])
        if rows:
            _add_styled_table(
                doc,
                ["Risk", "Likelihood", "Impact", "Mitigation"],
                rows,
            )
        else:
            doc.add_paragraph(risks)
    else:
        doc.add_paragraph("No risks identified.")

    # ========================================================================
    # 12. COST-BENEFIT ANALYSIS
    # ========================================================================
    doc.add_heading("12. Cost-Benefit Analysis", level=1)

    doc.add_heading("12.1 Estimated Costs", level=2)
    doc.add_paragraph(data.get("estimated_costs", ""))

    doc.add_heading("12.2 Expected Benefits", level=2)
    doc.add_paragraph(data.get("expected_benefits", ""))

    # ========================================================================
    # 13. IMPLEMENTATION TIMELINE
    # ========================================================================
    doc.add_heading("13. Implementation Timeline", level=1)
    timeline = data.get("timeline", "")
    if timeline.strip():
        rows = []
        for line in timeline.strip().split("\n"):
            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 3:
                rows.append(parts[:3])
            elif line.strip():
                rows.append([line.strip(), "", ""])
        if rows:
            _add_styled_table(doc, ["Phase", "Duration", "Deliverables"], rows)
        else:
            doc.add_paragraph(timeline)
    else:
        doc.add_paragraph("No timeline specified.")

    # ========================================================================
    # 14. TESTING STRATEGY
    # ========================================================================
    doc.add_heading("14. Testing Strategy", level=1)
    doc.add_paragraph(data.get("testing_strategy", ""))

    # ========================================================================
    # 15. ASSUMPTIONS & CONSTRAINTS
    # ========================================================================
    doc.add_heading("15. Assumptions & Constraints", level=1)

    doc.add_heading("15.1 Assumptions", level=2)
    assumptions = data.get("assumptions", "")
    if assumptions.strip():
        for line in assumptions.strip().split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")
    else:
        doc.add_paragraph("No assumptions documented.")

    doc.add_heading("15.2 Constraints", level=2)
    constraints = data.get("constraints", "")
    if constraints.strip():
        for line in constraints.strip().split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")
    else:
        doc.add_paragraph("No constraints documented.")

    # ========================================================================
    # 16. RECOMMENDATIONS
    # ========================================================================
    doc.add_heading("16. Recommendations", level=1)
    doc.add_paragraph(data.get("recommendations", ""))

    # ========================================================================
    # 17. APPENDICES
    # ========================================================================
    doc.add_heading("17. Appendices", level=1)
    doc.add_paragraph(data.get("appendices", ""))

    # ========================================================================
    # FOOTER
    # ========================================================================
    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run(
        f"{data.get('project_name', 'Systems Analysis')} — "
        f"{data.get('organization', '')} — "
        f"v{data.get('version', '1.0')}"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # -- Save to buffer --
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _add_requirements_list(doc, text: str):
    """Parse requirements text (one per line) into a numbered list."""
    for i, line in enumerate(text.strip().split("\n"), 1):
        if line.strip():
            doc.add_paragraph(line.strip(), style="List Number")
