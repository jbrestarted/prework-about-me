"""
Reusable python-docx helpers for the CVN C5IMP Administrative Tool Suite.

Provides consistent styles, headers, footers, document control blocks,
revision history tables, approval tables, and formatted section content
for every Word template in the suite.

All content is UNCLASSIFIED template/placeholder material only.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- Palette (Navy staff styling) ---
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT = "D9E2F3"  # table header shading hex (no #)

DISTRO = ("DISTRIBUTION STATEMENT PLACEHOLDER - Insert appropriate "
          "distribution statement prior to release.")
DRAFT_BANNER = "DRAFT TEMPLATE - NOT OFFICIAL POLICY - UNCLASSIFIED"


# ----------------------------------------------------------------------
# Low-level XML helpers
# ----------------------------------------------------------------------
def _set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _add_page_number_field(paragraph):
    """Insert PAGE field into a paragraph (for footers)."""
    run = paragraph.add_run()
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)


def _set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


# ----------------------------------------------------------------------
# Document scaffolding
# ----------------------------------------------------------------------
def style_base(doc):
    """Apply consistent base styles across the suite."""
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in (
        ("Heading 1", 15, NAVY),
        ("Heading 2", 13, NAVY),
        ("Heading 3", 11.5, GRAY),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.color.rgb = NAVY
    title.font.bold = True


def add_header_footer(doc, doc_title, doc_number):
    """Standard header (title / banner) and footer (control no. / page)."""
    section = doc.sections[0]

    hdr = section.header.paragraphs[0]
    hdr.text = ""
    r1 = hdr.add_run(f"{doc_title}\n")
    r1.font.size = Pt(9)
    r1.font.bold = True
    r1.font.color.rgb = NAVY
    r2 = hdr.add_run(DRAFT_BANNER)
    r2.font.size = Pt(8)
    r2.font.color.rgb = GRAY
    hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER

    ftr = section.footer.paragraphs[0]
    ftr.text = ""
    fr = ftr.add_run(f"{doc_number}  |  Page ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = GRAY
    _add_page_number_field(ftr)
    fr2 = ftr.add_run(f"  |  {DISTRO}")
    fr2.font.size = Pt(7)
    fr2.font.color.rgb = GRAY
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_title_page(doc, title, subtitle, doc_number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\n\nAIRCRAFT CARRIER C5I MODERNIZATION PROGRAM (C5IMP)")
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = GOLD

    t = doc.add_paragraph(style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.add_run(title)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run(subtitle)
    sr.font.size = Pt(12)
    sr.font.italic = True
    sr.font.color.rgb = GRAY

    add_doc_control_table(doc, doc_number)

    d = doc.add_paragraph()
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dr = d.add_run("\n" + DRAFT_BANNER + "\n" + DISTRO)
    dr.font.size = Pt(9)
    dr.font.bold = True
    dr.font.color.rgb = GRAY

    doc.add_page_break()


def make_table(doc, headers, rows, widths=None):
    """Styled table with shaded, repeating header row."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_row = table.rows[0]
    _set_repeat_header(hdr_row)
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = NAVY
        _set_cell_shading(cell, LIGHT)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    return table


def add_doc_control_table(doc, doc_number):
    doc.add_paragraph()
    make_table(
        doc,
        ["Document Control Field", "Entry"],
        [
            ["Document Number", doc_number],
            ["Version", "0.1 (DRAFT)"],
            ["Effective Date", "[DD MMM YYYY]"],
            ["Prepared By", "[Action Officer Name / Code]"],
            ["Reviewed By", "[Branch Head / Code]"],
            ["Approved By", "[Approval Authority / Code]"],
            ["Next Review Date", "[DD MMM YYYY]"],
            ["Classification", "UNCLASSIFIED (template) - "
             "[Insert handling caveat placeholder]"],
            ["SharePoint Location", "[C5IMP Site / Library / Folder]"],
        ],
        widths=[2.2, 4.3],
    )


def add_revision_history(doc):
    doc.add_heading("Revision History", level=1)
    make_table(
        doc,
        ["Version", "Date", "Author / Code", "Description of Change",
         "Approved By"],
        [
            ["0.1", "[DD MMM YYYY]", "[Name / Code]", "Initial draft template.",
             "[Name / Code]"],
            ["", "", "", "", ""],
            ["", "", "", "", ""],
        ],
    )


def add_approval_table(doc, roles=None):
    doc.add_heading("Approval", level=1)
    roles = roles or [
        "TYCOM N6 (or designated representative)",
        "NAVSEA Carrier Modernization Lead [Code Placeholder]",
        "PEO Carriers Representative [Code Placeholder]",
        "PEO C4I / PMW Representative [Code Placeholder]",
    ]
    make_table(
        doc,
        ["Role / Organization", "Name", "Signature", "Date"],
        [[r, "[Name]", "", "[DD MMM YYYY]"] for r in roles],
        widths=[2.8, 1.6, 1.3, 0.9],
    )


def add_section(doc, heading, paragraphs=(), bullets=(), level=1):
    doc.add_heading(heading, level=level)
    for para in paragraphs:
        doc.add_paragraph(para)
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")


def new_document(title, subtitle, doc_number):
    """Create a document with the full standard scaffold."""
    doc = Document()
    style_base(doc)
    add_header_footer(doc, title, doc_number)
    add_title_page(doc, title, subtitle, doc_number)
    add_revision_history(doc)
    return doc


def finish_document(doc, path, approval_roles=None):
    add_approval_table(doc, approval_roles)
    doc.save(path)
    return path
