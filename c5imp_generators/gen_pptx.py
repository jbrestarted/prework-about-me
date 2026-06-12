"""
Generates CVN_C5IMP_SharePoint_Design_Brief.pptx (16 slides) into
03_SharePoint_Design_Brief.

Usage: python gen_pptx.py <output_root>
"""

import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BANNER = "DRAFT - UNCLASSIFIED - TEMPLATE / PLACEHOLDER CONTENT ONLY"


def _style_title(shape, size=30, color=NAVY):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = color
            r.font.name = "Calibri"


def _add_banner(slide, prs):
    box = slide.shapes.add_textbox(Inches(0.3), prs.slide_height -
                                   Inches(0.45), prs.slide_width -
                                   Inches(0.6), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    p.text = BANNER
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = GRAY


def add_bullet_slide(prs, title, bullets, notes=None):
    """bullets: list of (level, text) or plain strings (level 0)."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    _style_title(slide.shapes.title)
    body = slide.placeholders[1].text_frame
    body.clear()
    first = True
    for b in bullets:
        level, text = b if isinstance(b, tuple) else (0, b)
        p = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        p.text = text
        p.level = level
        for r in p.runs:
            r.font.size = Pt(20 - 3 * level)
            r.font.name = "Calibri"
            if level == 0:
                r.font.color.rgb = NAVY
    _add_banner(slide, prs)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_table_slide(prs, title, headers, rows, col_widths=None,
                    font_size=11):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    _style_title(slide.shapes.title, size=26)
    nrows, ncols = len(rows) + 1, len(headers)
    left, top = Inches(0.4), Inches(1.4)
    width = prs.slide_width - Inches(0.8)
    height = prs.slide_height - Inches(2.1)
    table = slide.shapes.add_table(nrows, ncols, left, top, width,
                                   height).table
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            table.columns[i].width = int(width * w / total)
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(font_size)
                r.font.bold = True
                r.font.color.rgb = WHITE
    for ri, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            cell = table.cell(ri, c)
            cell.text = str(v)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size - 1)
    _add_banner(slide, prs)
    return slide


def build(path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = ("CVN C5IMP Modernization SharePoint "
                           "Administrative Tool Suite")
    _style_title(s.shapes.title, size=36)
    sub = s.placeholders[1]
    sub.text = ("SharePoint Design Brief for Leadership and Site-Owner "
                "Approval\nAircraft Carrier C5I Modernization Program "
                "(C5IMP)\n[Prepared by: Name / Code]  |  [DD MMM YYYY]")
    for p in sub.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(16)
            r.font.color.rgb = GRAY
    _add_banner(s, prs)

    # Slide 2: Purpose
    add_bullet_slide(prs, "Purpose", [
        "Establish a single SharePoint site as the authoritative "
        "administrative backbone for CVN C5I modernization.",
        (1, "Centralized modernization control: one master tracker, one "
            "baseline register, one action list."),
        (1, "Governance discipline: battle rhythm, agendas, minutes, and "
            "CCB packages in one controlled location."),
        (1, "Decision traceability: every decision indexed, linked to its "
            "record, rationale, and follow-on actions."),
        (1, "Readiness alignment: assessment results visible to TYCOM and "
            "Ship's Force at each milestone gate."),
        (1, "NAVSEA/TYCOM coordination: shared, permissioned visibility "
            "across acquisition and fleet stakeholders.")])

    # Slide 3: User Communities
    add_table_slide(prs, "User Communities",
        ["Community", "Primary Need", "Typical Access"],
        [["TYCOM N6 / N43 / N4", "Readiness governance, prioritization",
          "Edit"],
         ["NAVSEA Carrier Modernization", "Integration engineering, "
          "alteration management", "Edit"],
         ["PEO Carriers", "Platform acquisition coordination",
          "Contribute"],
         ["PEO C4I / PMW", "System fielding, ILS, install requirements",
          "Contribute"],
         ["Ship's Force C5I", "Impact awareness, readiness inputs",
          "Read (+ scoped input)"],
         ["RMC / Shipyard / Installation teams", "Availability "
          "execution data", "Contribute (scoped)"],
         ["Fleet staff [Placeholder]", "Schedule alignment", "Read"],
         ["Program office action officers", "Trackers, intake, actions",
          "Edit"]],
        col_widths=[3, 5, 2], font_size=13)

    # Slide 4: Operational Problem
    add_bullet_slide(prs, "Operational Problem - Current Administrative "
                          "Friction", [
        "Scattered files: program artifacts spread across email, shared "
        "drives, and personal folders.",
        "Inconsistent trackers: multiple competing spreadsheets with "
        "conflicting status for the same items.",
        "Poor decision traceability: decisions buried in minutes and "
        "email threads; rationale and conditions lost.",
        "Weak configuration visibility: installed versus documented "
        "baseline unclear by hull.",
        "Unclear ownership: actions assigned verbally, not tracked to "
        "closure.",
        "Fragmented modernization status: every status request triggers "
        "a manual data call to multiple organizations."])

    # Slide 5: Information Architecture
    add_table_slide(prs, "Proposed SharePoint Information Architecture",
        ["Page", "Function"],
        [["Home", "Orientation, news, battle rhythm, key POCs"],
         ["Modernization Dashboard", "Status rollups, hot items, "
          "future Power BI placeholder"],
         ["Intake and Change Review", "Submit and track proposed "
          "changes"],
         ["Configuration Baseline", "Approved baseline and change "
          "posture by hull/system"],
         ["Availability Integration", "Work packages, access, testing, "
          "turnover"],
         ["Risks / Issues / Actions", "Registers and burn-down views"],
         ["Governance and Meetings", "Calendar, agendas, minutes"],
         ["Decisions and CCB", "Decision log, records, CCB packages "
          "(restricted)"],
         ["Readiness Assessments", "Hull/system readiness results"],
         ["References and Templates", "All controlled templates and "
          "references"],
         ["Lessons Learned", "After Action Reports and search"]],
        col_widths=[3, 7], font_size=12)

    # Slide 6: Libraries and Lists
    add_table_slide(prs, "Recommended Libraries and Lists",
        ["Library / List", "Holds", "Control Features"],
        [["Program Documents Library", "Charters, SOPs, plans",
          "Major versions, check-out, approval"],
         ["Templates Library", "Blank templates only", "Read-only for "
          "users; owner-managed"],
         ["Modernization Tracker Library", "Authoritative Excel "
          "trackers", "Browser editing; version history"],
         ["Configuration Baseline Library", "Baseline register, "
          "references", "CM-staff edit only"],
         ["Decision Records Library", "Completed decision records",
          "Unique permissions (restricted group)"],
         ["Meeting Records Library", "Agendas, read-aheads, minutes",
          "Filed by series and date"],
         ["Risk and Action Lists", "Risk and action items for web "
          "part views", "Filtered views, conditional formatting"],
         ["Stakeholder Directory List", "POC roster", "Feeds People "
          "web parts"],
         ["Lessons Learned Library", "After Action Reports",
          "Metadata-driven search"]],
        col_widths=[3.2, 3.6, 3.4], font_size=12)

    # Slide 7: Metadata Strategy
    add_bullet_slide(prs, "Metadata Strategy", [
        "Site columns applied consistently across libraries and lists - "
        "metadata drives every view and dashboard.",
        (1, "Hull  |  System Category  |  Program Office  |  Document "
            "Type  |  Status"),
        (1, "Decision Authority  |  Availability  |  Fiscal Year  |  "
            "Owner  |  Review Date  |  Last Updated"),
        (1, "CCB Required  |  Readiness Impact  |  Classification/"
            "Handling Caveat [placeholder - tailor per command "
            "guidance]"),
        "Required columns enforced at upload; monthly metadata audit by "
        "site owners.",
        "Folders are minimized - filtering and grouping come from "
        "metadata views, not folder trees."])

    # Slide 8: Recommended Web Parts
    add_table_slide(prs, "Recommended Web Parts",
        ["Web Part", "Employment"],
        [["Hero", "Major navigation tiles on Home"],
         ["Quick Links", "Key workflows: submit intake, open trackers, "
          "get templates"],
         ["Document Library", "Controlled document access via named "
          "views"],
         ["List", "Action items, risks, and decision log views"],
         ["Highlighted Content", "Recently updated and high-priority "
          "files, automatic by metadata"],
         ["News", "Modernization updates and announcements"],
         ["Events", "Battle rhythm and CCB calendar"],
         ["People", "Key POCs by page"],
         ["Text", "Short process guidance blocks"],
         ["Embed / Power BI", "Placeholder for future dashboard "
          "integration (tenant approval required)"],
         ["Forms placeholder", "Intake workflow if tenant policy "
          "allows; fallback is list new-item form"]],
        col_widths=[3, 7], font_size=12)

    # Slide 9: Page-by-page employment plan
    add_table_slide(prs, "Page-by-Page Web Part Employment Plan",
        ["Page", "Web Parts", "Content Sources"],
        [["Home", "Hero, News, Events, Quick Links, People",
          "News; calendar; directory"],
         ["Modernization Dashboard", "Highlighted Content, List x2, "
          "Embed placeholder", "Trackers; risk/action lists"],
         ["Intake and Change Review", "Quick Links, List, Text, Forms "
          "placeholder", "Intake list; templates"],
         ["Configuration Baseline", "Document Library, List, Text",
          "Baseline library"],
         ["Availability Integration", "Document Library, List, Events",
          "Availability tracker; calendar"],
         ["Risks / Issues / Actions", "List x3, Highlighted Content",
          "Risk and action lists"],
         ["Governance and Meetings", "Events, Document Library, Quick "
          "Links", "Meeting records; calendar"],
         ["Decisions and CCB", "List, Document Library, Text",
          "Decision records (restricted)"],
         ["Readiness Assessments", "Document Library, Highlighted "
          "Content", "Readiness products"],
         ["References and Templates", "Document Library, Quick Links, "
          "Text", "Templates library"],
         ["Lessons Learned", "Document Library, Highlighted Content",
          "Lessons Learned library"]],
        col_widths=[2.8, 4.2, 3.0], font_size=11)

    # Slide 10: Tool suite
    add_table_slide(prs, "Administrative Tool Suite Hosted on the Site",
        ["Word Templates (Program Documents / Templates)",
         "Excel Tools (Tracker Library)"],
        [["Program Charter; Governance SOP; CCB Charter",
          "Master Modernization Tracker"],
         ["Intake Form; Change Evaluation Memorandum",
          "Baseline Configuration Register"],
         ["Modernization Readiness Assessment",
          "Risk / Issue / Opportunity Register"],
         ["Availability Integration Planning Guide",
          "Action Item Tracker; Decision Log"],
         ["Stakeholder Engagement Plan",
          "Stakeholder RACI Matrix"],
         ["Meeting Minutes; Decision Record; AAR templates",
          "Availability Integration Tracker"],
         ["SharePoint User Guide",
          "Readiness Assessment Workbook; Battle Rhythm Tracker; "
          "Content Inventory & Metadata Plan"]],
        col_widths=[5, 5], font_size=12)

    # Slide 11: Workflow model
    add_bullet_slide(prs, "Workflow Model - Modernization Lifecycle", [
        "Intake  >  Triage  >  Stakeholder Review  >  CCB / Decision  >  "
        "Schedule Integration  >  Execution  >  Readiness Assessment  >  "
        "Lessons Learned  >  Baseline Update",
        (1, "Intake: sponsor submits via Intake page; tracking number "
            "assigned."),
        (1, "Triage and review: WIPT coordinates the multi-discipline "
            "Change Evaluation Memorandum."),
        (1, "Decision: CCB dispositions; Decision Record filed and "
            "logged."),
        (1, "Integration and execution: availability tracker and master "
            "tracker drive status."),
        (1, "Closure: readiness assessment gates turnover; AAR feeds "
            "process improvement; baseline register updated."),
        "Every stage writes to a SharePoint artifact - the workflow is "
        "auditable end to end."])

    # Slide 12: Governance model
    add_bullet_slide(prs, "Governance Model - How TYCOM and NAVSEA Use "
                          "the Site", [
        "TYCOM: readiness visibility and decision control.",
        (1, "Readiness Assessments and Dashboard pages support milestone "
            "gates and fleet alignment."),
        (1, "CCB membership and decision log give N6/N43 traceable "
            "configuration governance."),
        "NAVSEA: integration engineering and configuration alignment.",
        (1, "Baseline register and availability tracker keep alteration "
            "and installation posture current."),
        (1, "Work package views synchronize NAVSEA, RMC, and shipyard "
            "execution."),
        "Shared: one battle rhythm, one action tracker, one decision "
        "log - no parallel bookkeeping.",
        "Chairs enforce the rule: if it is not on the site, it is not "
        "the record."])

    # Slide 13: Permissions
    add_table_slide(prs, "Permissions and Information Control",
        ["Group", "Access", "Notes"],
        [["Site Owners", "Full Control", "2-3 named owners; structure "
          "and permissions"],
         ["TYCOM Editors", "Edit", "All content except restricted "
          "records"],
         ["NAVSEA Editors", "Edit", "All content except restricted "
          "records"],
         ["PEO/PMW Contributors", "Contribute", "Trackers, intake, "
          "meeting records"],
         ["Ship Viewers", "Read", "Unrestricted content"],
         ["Shipyard/RMC Contributors", "Contribute (scoped)",
          "Availability content only"],
         ["Read-Only Leadership", "Read", "Executive visibility"],
         ["Restricted Decision Records Group", "Edit (unique "
          "permissions)", "Decision Records library only"]],
        col_widths=[3.5, 2.5, 4], font_size=13)

    # Slide 14: Roadmap
    add_table_slide(prs, "Implementation Roadmap",
        ["Phase", "Scope", "Target (Placeholder)"],
        [["1. Build site shell", "Pages, navigation, permission groups",
          "[Week 1-2]"],
         ["2. Create libraries/lists", "Per Content Inventory workbook",
          "[Week 2-3]"],
         ["3. Apply metadata", "Site columns, required fields",
          "[Week 3-4]"],
         ["4. Upload templates", "Full Word/Excel suite", "[Week 4]"],
         ["5. Configure views", "By Hull, Overdue, Pending Decision, "
          "etc.", "[Week 5]"],
         ["6. Add web parts", "Per Web Part Employment Matrix",
          "[Week 5-6]"],
         ["7. Pilot", "One hull or one availability end-to-end",
          "[Week 7-18]"],
         ["8. Refine governance", "SOP/CCB updates from pilot AAR",
          "[Week 18-20]"],
         ["9. Expand to fleet-level use", "Onboard remaining hulls; "
          "training", "[Week 20+]"]],
        col_widths=[3, 5, 2], font_size=13)

    # Slide 15: Success criteria
    add_bullet_slide(prs, "Success Criteria", [
        "Reduced duplicate trackers: one authoritative tracker per "
        "function; retired spreadsheets archived.",
        "Improved decision traceability: 100% of decisions logged with "
        "record, authority, and follow-on actions.",
        "Faster modernization status reporting: monthly report built "
        "from live views, not data calls.",
        "Clear action ownership: zero unowned actions; overdue rate "
        "visible and trending down.",
        "Better configuration visibility: installed-vs-documented "
        "baseline current for every hull.",
        "Improved readiness assessment consistency: common criteria "
        "bank used at every gate.",
        "Stronger NAVSEA/TYCOM coordination: shared battle rhythm and "
        "single source of record."])

    # Slide 16: Next steps
    add_bullet_slide(prs, "Recommended Next Steps", [
        "1. Approve this design brief and designate 2-3 site owners.",
        "2. Confirm tenant policy on Forms and Power BI web parts.",
        "3. Stand up the site shell and permission groups (Phase 1).",
        "4. Build libraries, lists, and metadata per the Content "
        "Inventory and Metadata Plan workbook (Phases 2-3).",
        "5. Upload the administrative tool suite and configure views "
        "(Phases 4-5).",
        "6. Assemble pages per the Web Part Employment Matrix "
        "(Phase 6).",
        "7. Select the pilot: [CVN-XX] [PIA FY2X placeholder] entering "
        "its planning window; run one full intake-to-decision cycle.",
        "8. Hold a pilot AAR at [pilot + 90 days]; refine and expand "
        "fleet-wide."])

    prs.save(path)
    print(f"  [pptx] {path}")


def build_all(root):
    d = os.path.join(root, "03_SharePoint_Design_Brief")
    os.makedirs(d, exist_ok=True)
    build(os.path.join(d, "CVN_C5IMP_SharePoint_Design_Brief.pptx"))


if __name__ == "__main__":
    build_all(sys.argv[1] if len(sys.argv) > 1 else
              "CVN_C5IMP_Admin_Tool_Suite")
