"""
Reusable openpyxl helpers for the CVN C5IMP Administrative Tool Suite.

Provides consistent table styling, freeze panes, data-validation dropdowns
backed by named ranges, conditional formatting, and instructions tabs.
"""

import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.defined_name import DefinedName

NAVY_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10, name="Calibri")
BODY_FONT = Font(size=10, name="Calibri")
TITLE_FONT = Font(size=14, bold=True, color="1F3A5F", name="Calibri")
NOTE_FONT = Font(size=10, italic=True, color="595959", name="Calibri")

RED_FILL = PatternFill("solid", fgColor="F8CBAD")
YELLOW_FILL = PatternFill("solid", fgColor="FFE699")
GREEN_FILL = PatternFill("solid", fgColor="C6E0B4")
ORANGE_FILL = PatternFill("solid", fgColor="F4B183")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ----------------------------------------------------------------------
# Standard controlled-vocabulary lists (placeholder values only)
# ----------------------------------------------------------------------
LISTS = {
    "StatusList": ["Not Started", "In Progress", "On Hold", "Blocked",
                   "Pending Decision", "Complete", "Cancelled"],
    "PriorityList": ["Critical", "High", "Medium", "Low"],
    "RiskLevelList": ["High", "Moderate", "Low"],
    "HullList": ["CVN-XX", "CVN-YY", "CVN-ZZ", "Class-Wide"],
    "SystemCategoryList": ["C2 Suite", "Communications", "Networks",
                           "Combat Systems I/F", "ISR", "Navigation",
                           "Cyber Infrastructure"],
    "OrgList": ["TYCOM N6", "TYCOM N43", "TYCOM N4", "NAVSEA",
                "PEO Carriers", "PEO C4I/PMW", "Ship's Force", "RMC",
                "Shipyard", "Program Office", "Fleet Staff"],
    "DecisionStatusList": ["Submitted", "In Evaluation", "Pending Decision",
                           "Approved", "Approved w/ Conditions",
                           "Disapproved", "Deferred"],
    "DocTypeList": ["Charter", "SOP", "Form", "Memorandum", "Assessment",
                    "Plan", "Minutes", "Decision Record", "AAR", "Guide",
                    "Tracker"],
    "AvailTypeList": ["PIA", "DPIA", "CIA", "Emergent Window"],
    "RatingList": ["Green", "Yellow", "Red", "Gray"],
    "YesNoList": ["Yes", "No"],
    "RACIList": ["R", "A", "C", "I", "R/A", ""],
    "ScoreList": ["1", "2", "3", "4", "5"],
}


def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    _add_lookups(wb)
    return wb


def _add_lookups(wb):
    """Hidden Lookups sheet backing named ranges for all dropdowns."""
    ws = wb.create_sheet("Lookups")
    for col, (name, values) in enumerate(LISTS.items(), start=1):
        letter = get_column_letter(col)
        ws.cell(row=1, column=col, value=name).font = Font(bold=True, size=9)
        for r, v in enumerate(values, start=2):
            ws.cell(row=r, column=col, value=v)
        ref = f"Lookups!${letter}$2:${letter}${len(values) + 1}"
        wb.defined_names.add(DefinedName(name, attr_text=ref))
    ws.sheet_state = "hidden"


def _table_name(ws_title):
    name = "tbl" + re.sub(r"[^A-Za-z0-9]", "", ws_title)
    return name[:30] or "tblData"


def add_table_sheet(wb, title, columns, rows, widths=None, validations=None,
                    cond_rules=None, extra_blank_rows=15):
    """
    Create a sheet with a styled Excel table.

    columns: list of header strings.
    rows: list of row tuples/lists (example placeholder data).
    widths: optional list of column widths.
    validations: {header_name: named_range_name} dropdown mapping.
    cond_rules: list of (header_name, match_text, fill) tuples applying
                'cell equals text' highlighting to that column.
    extra_blank_rows: blank in-table rows left for users to fill.
    """
    ws = wb.create_sheet(title[:31])
    ncols = len(columns)

    for c, h in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HDR_FONT
        cell.fill = NAVY_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER

    nrows = len(rows) + extra_blank_rows
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
    for r in range(len(rows) + 2, nrows + 2):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).border = BORDER

    last_col = get_column_letter(ncols)
    table = Table(displayName=_table_name(title),
                  ref=f"A1:{last_col}{nrows + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2",
                                          showRowStripes=True)
    ws.add_table(table)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = None  # table provides its own filter buttons

    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    else:
        for i in range(1, ncols + 1):
            ws.column_dimensions[get_column_letter(i)].width = 22

    col_index = {h: i for i, h in enumerate(columns, start=1)}

    for header, named_range in (validations or {}).items():
        if header not in col_index:
            continue
        letter = get_column_letter(col_index[header])
        dv = DataValidation(type="list", formula1=f"={named_range}",
                            allow_blank=True, showDropDown=False)
        dv.error = "Select a value from the dropdown list."
        dv.errorTitle = "Invalid entry"
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{nrows + 1}")

    for header, text, fill in (cond_rules or []):
        if header not in col_index:
            continue
        letter = get_column_letter(col_index[header])
        rng = f"{letter}2:{letter}{nrows + 1}"
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=[f'"{text}"'],
                            fill=fill))
    return ws


def add_overdue_rule(ws, date_header, status_header, columns, nrows):
    """Red-fill rows whose due date passed and status is not Complete."""
    col_index = {h: i for i, h in enumerate(columns, start=1)}
    if date_header not in col_index or status_header not in col_index:
        return
    d = get_column_letter(col_index[date_header])
    s = get_column_letter(col_index[status_header])
    last = get_column_letter(len(columns))
    rng = f"A2:{last}{nrows + 1}"
    formula = (f'AND(ISNUMBER(${d}2),${d}2<TODAY(),'
               f'${s}2<>"Complete",${s}2<>"Cancelled")')
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula],
                                                   fill=RED_FILL))


def add_instructions_sheet(wb, workbook_title, purpose, steps,
                           maintenance=None, first=True):
    ws = wb.create_sheet("Instructions", 0 if first else None)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110

    ws["A1"] = workbook_title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("CVN C5I Modernization Program (C5IMP) - "
                "DRAFT TEMPLATE - NOT OFFICIAL POLICY - UNCLASSIFIED. "
                "Placeholder data only.")
    ws["A2"].font = NOTE_FONT

    r = 4
    ws.cell(row=r, column=1, value="Purpose").font = Font(bold=True, size=11)
    r += 1
    ws.cell(row=r, column=1, value=purpose).alignment = \
        Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.row_dimensions[r].height = 45
    r += 2

    ws.cell(row=r, column=1, value="How to Use This Workbook").font = \
        Font(bold=True, size=11)
    r += 1
    for s in steps:
        c = ws.cell(row=r, column=1, value="- " + s)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = BODY_FONT
        ws.row_dimensions[r].height = max(15, 15 * (len(s) // 100 + 1))
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Maintenance Rules").font = \
        Font(bold=True, size=11)
    r += 1
    for m in (maintenance or [
        "The SharePoint copy of this workbook is authoritative. Edit in "
        "the browser; do not download/re-upload working copies.",
        "Update within 2 business days of any governance event per the "
        "C5IMP Governance SOP.",
        "Use dropdown values only in controlled columns; do not overwrite "
        "data validation or table formatting.",
        "Do not delete rows with history; mark items Complete or "
        "Cancelled instead.",
        "Do not enter classified information; use approved placeholders "
        "and references."]):
        c = ws.cell(row=r, column=1, value="- " + m)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = BODY_FONT
        r += 1
    return ws
